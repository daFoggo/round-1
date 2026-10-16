import os
import re
from datetime import datetime
import openpyxl
import pandas as pd
from typing import Dict, List, Tuple, Any, Optional

def clean_text(val: Any) -> str:
    """Làm sạch chuỗi, xóa footnotes như (1), 2/, khoảng trắng thừa."""
    if val is None:
        return ""
    s = str(val).strip().replace('\r', ' ').replace('\n', ' ')
    s = re.sub(r'\(\d+\)', '', s)
    s = re.sub(r'(?<=\S)\([a-z]\)', '', s)
    s = re.sub(r'\b\d+/(?=\s|$)', '', s)
    s = re.sub(r'\s+', ' ', s)
    return s.strip()

_PLACEHOLDER_RE = re.compile(r'^(?:[-–—.]+|n/?a|x|z|nm|\*+|\(-?\d{1,3}-?\)|\([a-z]\)|\(\w{1,2}\))$', re.IGNORECASE)
_NUMBER_RE = re.compile(r'^\(?-?[$€£]?\s*-?\d[\d,]*\.?\d*\s*%?\)?$|^\(?-?[$€£]?\s*-?\.\d+\s*%?\)?$')


def is_placeholder(v: Any) -> bool:
    """Ô giữ chỗ không mang giá trị: '-', em dash, '..', 'n/a', '(i)', '(-8-)'."""
    return isinstance(v, str) and bool(_PLACEHOLDER_RE.match(v.strip().lstrip('$€£').strip()))


def parse_numeric(v: Any) -> Optional[float]:
    """Chuyển '$1,402', '50%', '(123)' (kế toán, số âm) thành số; không phải số -> None."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    if not isinstance(v, str):
        return None
    t = v.strip()
    if not t or not _NUMBER_RE.match(t):
        return None
    neg = t.startswith('(') and t.endswith(')')
    t = re.sub(r'[()$€£%,\s]', '', t)
    try:
        num = float(t) if '.' in t else int(t)
    except ValueError:
        return None
    return -abs(num) if neg else num


def is_num_or_date(v: Any) -> bool:
    if v is None:
        return False
    if isinstance(v, (int, float, datetime)):
        return True
    return parse_numeric(v) is not None


def clean_label(val: Any) -> str:
    """Nhãn dòng: bỏ chú thích dính chữ cuối nhãn (vd. 'Private industry3' -> 'Private industry')."""
    s = clean_text(val)
    return re.sub(r'(?<=[a-z)])\d(?:,\d)*$', '', s).strip()


def indent_width(ws, r: int, c: int, raw: Any) -> int:
    """Độ thụt lề của nhãn: khoảng trắng đầu dòng + alignment.indent của Excel."""
    lead = 0
    if isinstance(raw, str):
        lead = len(raw.replace('	', '    ')) - len(raw.replace('	', '    ').lstrip())
    try:
        lead += int(ws.cell(r, c).alignment.indent or 0) * 2
    except Exception:
        pass
    return lead


def is_note_row(c1_val: str) -> bool:
    """Nhận diện các dòng chú thích cuối bảng (Footnotes, Notes)."""
    c1_low = c1_val.lower().strip()
    return (
        c1_low.startswith('footnote') or
        c1_low.startswith('note:') or
        bool(re.match(r'^\(\d+\)\s*', c1_val)) or
        bool(re.match(r'^\d{1,2}/(?:\s|$)', c1_val))  # '1/ Ghi chú' (không phải '2005/2006')
    )

_CODE_RE = re.compile(r'^[A-Z]{1,3}[A-Z0-9]{1,5}$')


def _is_code_row(first_val: str, row_info, idx: int) -> bool:
    """Dòng chỉ gồm mã cột ngắn viết hoa (DJ54, DJ57...), thường nằm dưới một tiêu đề khối."""
    return bool(first_val) and bool(_CODE_RE.match(first_val)) and row_info[idx]['nums'] == 0


def detect_table_ranges(ws: openpyxl.worksheet.worksheet.Worksheet, grid: Dict[Tuple[int, int], Any], max_row: int, max_col: int) -> List[Tuple[int, int, str]]:
    """
    Phát hiện các khối bảng độc lập trên cùng một worksheet (ví dụ Fiscal years vs Calendar years).
    Trả về danh sách [(start_row, end_row, table_title), ...].
    """
    row_info = []
    for r in range(1, max_row + 1):
        vals = [grid.get((r, c)) for c in range(1, max_col + 1)]
        non_empty = [v for v in vals if v is not None and str(v).strip() != '']
        nums = sum(1 for v in non_empty if isinstance(v, (int, float)) or (isinstance(v, str) and parse_numeric(v) is not None))
        placeholders = sum(1 for v in non_empty if is_placeholder(v))
        years = sum(1 for v in non_empty if isinstance(v, int) and 1900 <= v <= 2100)
        c1 = str(vals[0]).strip() if vals and vals[0] is not None else ''
        is_note = is_note_row(c1)
        row_info.append({
            'r': r,
            'non_empty': len(non_empty),
            'nums': nums,
            'placeholders': placeholders,
            'years': years,
            'is_note': is_note,
            'c1': c1,
            'first_val': clean_text(non_empty[0]) if non_empty else ''
        })

    tables = []
    curr_start = 1
    had_data = False
    table_title = row_info[0]['first_val'] or "Table_1"

    for idx, info in enumerate(row_info):
        r = info['r']
        if info['is_note']:
            continue
        if (info['nums'] >= 2 or (info['nums'] >= 1 and info['placeholders'] >= 1)) and info['years'] < info['nums']:
            had_data = True
        elif had_data:
            # Phát hiện nếu có bảng mới xuất hiện sau một khối dữ liệu
            if info['years'] >= 2 or (info['non_empty'] >= 2 and info['nums'] == 0 and info['placeholders'] == 0 and not info['is_note']):
                next_has_years_or_data = False
                for look in row_info[idx+1:idx+4]:
                    if look['years'] >= 2 or look['nums'] >= 2:
                        next_has_years_or_data = True
                        break
                if next_has_years_or_data:
                    tables.append((curr_start, r - 1, table_title))
                    curr_start = r
                    table_title = info['first_val'] or f"Table_{len(tables)+1}"
                    # Dòng mã cột (DJ54...) không phải tên bảng: lấy nhãn văn bản ngay phía trên (vd. "Indices")
                    if _is_code_row(info['first_val'], row_info, idx):
                        for back in row_info[max(0, idx - 3):idx][::-1]:
                            if back['non_empty'] == 1 and back['nums'] == 0 and back['first_val'] and not back['is_note']:
                                table_title = back['first_val']
                                break
                    had_data = False

    tables.append((curr_start, max_row, table_title))
    return tables

def _row_numbers(grid: Dict[Tuple[int, int], Any], r: int, max_col: int) -> List[float]:
    vals = []
    for c in range(2, max_col + 1):
        v = grid.get((r, c))
        if is_num_or_date(v) and not isinstance(v, datetime):
            vals.append(parse_numeric(v))
    return vals


def _is_year_header_row(grid: Dict[Tuple[int, int], Any], r: int, max_col: int) -> bool:
    """Dòng mà mọi giá trị số đều là năm (1900-2100) và dòng kế tiếp có số liệu không phải năm."""
    vals = _row_numbers(grid, r, max_col)
    if len(vals) < 2 or not all(float(v).is_integer() and 1900 <= v <= 2100 for v in vals):
        return False
    for look in range(r + 1, r + 5):  # bỏ qua vài dòng trống/nhãn nhóm giữa header và dữ liệu
        nxt = _row_numbers(grid, look, max_col)
        if len(nxt) >= 2:
            return not all(float(v).is_integer() and 1900 <= v <= 2100 for v in nxt)
    return False


def _join_wrapped(parts: List[str]) -> List[str]:
    """Nối các mảnh tiêu đề bị ngắt dòng trong cùng một ô ('Rubber,' + 'plastics &' + 'non-metallic' ...) thành một cụm.
    Chỉ nối khi mảnh trước kết thúc bằng ',' '&' '/' hoặc mảnh sau bắt đầu bằng '&' / chữ thường; tên cha-con thật
    (viết hoa đầu) được giữ nguyên."""
    out: List[str] = []
    for part in parts:
        if out and (out[-1].endswith((',', '&', '/')) or part.startswith('&') or part[:1].islower()):
            joiner = '' if (out[-1].endswith('-') and part[:1].islower()) else ' '
            out[-1] = out[-1] + joiner + part
        else:
            out.append(part)
    return out


def parse_table_block(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    grid: Dict[Tuple[int, int], Any],
    start_row: int,
    end_row: int,
    max_col: int
) -> pd.DataFrame:
    """Parse một khối bảng cụ thể (từ start_row đến end_row) thành DataFrame phẳng."""
    # 1. Tìm data_start_row trong khối
    data_start = None
    for r in range(start_row, end_row + 1):
        c1 = grid.get((r, 1))
        if isinstance(c1, int) and c1 == 1:
            data_start = r
            break
        nums = sum(1 for c in range(2, max_col + 1) if is_num_or_date(grid.get((r, c))))
        if nums >= 2 and r < end_row and _is_year_header_row(grid, r, max_col):
            continue  # dòng chỉ gồm năm (2022 | 2023 ...) là header, dữ liệu bắt đầu ở dòng sau
        placeholders = sum(1 for c in range(2, max_col + 1) if is_placeholder(grid.get((r, c))))
        if nums >= 2 or (nums >= 1 and placeholders >= 1):
            r_next_strings = sum(1 for c in range(2, max_col + 1) if is_num_or_date(grid.get((r, c))) and isinstance(grid.get((r+1, c)), str) and len(str(grid.get((r+1, c)))) > 3 and parse_numeric(grid.get((r+1, c))) is None and not is_placeholder(grid.get((r+1, c))))
            if r_next_strings >= 2:
                continue
            data_start = r
            break

    if data_start is None:
        data_start = min(start_row + 1, end_row)

    # 2. Tìm header_end_row
    header_end = data_start - 1
    while header_end >= start_row:
        non_empty = [c for c in range(1, max_col + 1) if grid.get((header_end, c)) is not None and str(grid.get((header_end, c))).strip() != '']
        if len(non_empty) == 0 or (len(non_empty) == 1 and non_empty[0] == 1):
            header_end -= 1
        else:
            break

    if header_end < start_row:
        header_end = start_row

    # 3. Tìm header_start_row
    header_start = start_row
    for r in range(start_row, header_end + 1):
        raw_non_empty = [c for c in range(1, max_col + 1) if ws.cell(r, c).value is not None and str(ws.cell(r, c).value).strip() != '']
        if len(raw_non_empty) <= 1 and max_col > 2:
            txt = str(ws.cell(r, raw_non_empty[0]).value).strip().upper() if raw_non_empty else ''
            if any(k in txt for k in ['HOUSEHOLD DATA', 'ANNUAL AVERAGES', 'THOUSANDS', 'TABLE']) or len(txt) > 35:
                continue
        if len(raw_non_empty) == 0:
            continue
        header_start = r
        break

    # 4. Xây dựng Header Path cho từng cột
    col_names = []
    for c in range(1, max_col + 1):
        parts = []
        for r in range(header_start, header_end + 1):
            val = clean_text(grid.get((r, c)))
            if val and (not parts or parts[-1] != val):
                parts.append(val)
        col_names.append(" | ".join(_join_wrapped(parts)) if parts else f"Col_{c}")

    seen_cols = {}
    unique_col_names = []
    for c_idx, name in enumerate(col_names):
        name = name.strip() or f"Column_{c_idx+1}"
        if name in seen_cols:
            seen_cols[name] += 1
            unique_col_names.append(f"{name}_{seen_cols[name]}")
        else:
            seen_cols[name] = 1
            unique_col_names.append(name)

    # Kiểm tra xem Cột 1 có phải là Indent Level (như p4_t03) không
    c1_header = unique_col_names[0].lower() if unique_col_names else ""
    first_col_vals = [grid.get((r, 1)) for r in range(data_start, min(data_start + 5, end_row + 1))]
    looks_like_row_number = (
        len(first_col_vals) >= 3
        and all(isinstance(v, int) for v in first_col_vals)
        and all(b - a == 1 for a, b in zip(first_col_vals, first_col_vals[1:]))
    )  # cột 'No' = 1,2,3,4,5 là số thứ tự, không phải mức thụt lề
    is_indent_col = 'indent' in c1_header or (not looks_like_row_number and any(
        isinstance(v, int) and v in [0, 1, 2, 3, 4] for v in first_col_vals
    ))

    label_col_idx = 2 if is_indent_col and max_col >= 2 else 1

    # Cột 1 gần như trống (vd. nhãn năm/quý nằm ở cột 2-3): dùng các cột đầu có chữ/năm làm nhãn dòng
    label_cols = [label_col_idx]
    body_rows = [r for r in range(data_start, end_row + 1)
                 if any(grid.get((r, c)) not in (None, '') for c in range(1, max_col + 1))]
    if not is_indent_col and body_rows:
        def _filled(c):
            return sum(1 for r in body_rows if grid.get((r, c)) not in (None, '')) / len(body_rows)
        def _labelish(c):
            vals = [grid.get((r, c)) for r in body_rows if grid.get((r, c)) not in (None, '')]
            if not vals:
                return False
            ok = sum(1 for v in vals if (isinstance(v, str) and parse_numeric(v) is None)
                     or (isinstance(v, int) and 1900 <= v <= 2100))
            return ok / len(vals) >= 0.8
        if _filled(1) < 0.1:
            cand = []
            for c in range(2, min(max_col, 4) + 1):
                if _filled(c) >= 0.2 and _labelish(c):
                    cand.append(c)
                else:
                    break
            if cand:
                label_cols = cand
                label_col_idx = cand[-1]

    # Thụt lề của nhãn (khoảng trắng đầu dòng / alignment.indent) cho bảng không có cột Indent Level
    widths = {}
    use_indent = False
    if not is_indent_col and len(label_cols) == 1:
        for r in body_rows:
            raw = grid.get((r, label_col_idx))
            row_nums = sum(1 for v in [grid.get((r, c)) for c in range(label_col_idx + 1, max_col + 1)]
                           if is_num_or_date(v) or is_placeholder(v))
            if raw not in (None, '') and row_nums > 0:
                widths[r] = indent_width(ws, r, label_col_idx, raw)
        use_indent = len(set(widths.values())) >= 2
    ind_stack: List[Tuple[int, str]] = []
    last_year_label = ""

    # 5. Xây dựng Row_Path phân cấp
    data_rows = []
    major_section = ""
    sub_section = ""
    indent_stack: Dict[int, str] = {}

    for r in range(header_end + 1, data_start):
        s_val = clean_label(grid.get((r, label_col_idx)))
        if s_val and not is_note_row(s_val):
            major_section = s_val

    for r in range(data_start, end_row + 1):
        row_raw = [grid.get((r, c)) for c in range(1, max_col + 1)]
        if all(v is None or str(v).strip() == '' for v in row_raw):
            continue

        if len(label_cols) > 1:
            parts = [clean_label(row_raw[c - 1]) for c in label_cols]
            if parts[0]:
                last_year_label = parts[0]
            elif any(parts[1:]):
                parts[0] = last_year_label  # năm bị bỏ trống ở các dòng quý -> điền xuống
            raw_label = " | ".join(p for p in parts if p)
        else:
            raw_label = clean_label(row_raw[label_col_idx - 1])
            # Nhãn quý rút gọn ('Q3') kế thừa năm của dòng trước ('2010 Q2')
            if re.fullmatch(r'Q[1-4]', raw_label) and last_year_label:
                raw_label = f"{last_year_label} {raw_label}"
            ym = re.match(r'^(\d{4})\b', raw_label)
            if ym:
                last_year_label = ym.group(1)
        if is_note_row(raw_label):
            continue

        other_vals = row_raw[label_col_idx:]
        num_count = sum(1 for v in other_vals if is_num_or_date(v))

        # Trường hợp 1: Dòng thuần tiêu đề nhóm chính (nums == 0)
        placeholder_count = sum(1 for v in other_vals if is_placeholder(v))
        if raw_label != "" and num_count == 0 and placeholder_count == 0:
            major_section = raw_label
            sub_section = ""
            ind_stack = []
            continue

        # Trường hợp 2: Có cột Indent Level (cây phân cấp theo mức thụt lề số)
        if is_indent_col:
            lvl_val = row_raw[0]
            try:
                lvl = int(lvl_val) if lvl_val is not None else 0
            except (ValueError, TypeError):
                lvl = 0
            indent_stack[lvl] = raw_label
            indent_stack = {k: v for k, v in indent_stack.items() if k <= lvl}
            row_path = " | ".join(indent_stack[k] for k in sorted(indent_stack.keys()))
            current_sec = indent_stack.get(0, major_section)
        elif use_indent:
            w = widths.get(r, 0)
            while ind_stack and ind_stack[-1][0] >= w:
                ind_stack.pop()
            path_parts = [p for p in [major_section] + [lbl for _, lbl in ind_stack] + [raw_label] if p]
            row_path = " | ".join(path_parts)
            ind_stack.append((w, raw_label))
            current_sec = major_section
        else:
            # Group parent detection (BLS demographics + general hierarchical stubs):
            is_group_parent = (
                bool(re.search(r'(Total|Men|Women|Both sexes).*?,\s*\d+\s*years and over', raw_label, re.IGNORECASE)) or
                any(raw_label.lower().startswith(race) for race in ['white', 'black', 'asian', 'hispanic']) or
                raw_label.endswith(':') or
                any(raw_label.lower().startswith(kw) for kw in ['region:', 'sector:', 'department:', 'division:'])
            )

            if is_group_parent:
                sub_section = raw_label
                path_parts = [p for p in [major_section, sub_section] if p]
                row_path = " | ".join(path_parts)
            else:
                path_parts = [p for p in [major_section, sub_section, raw_label] if p]
                row_path = " | ".join(path_parts)
            current_sec = major_section

        # Làm sạch kiểu dữ liệu các ô
        cleaned_row = []
        for val in row_raw:
            if val is None or str(val).strip() == '':
                cleaned_row.append(None)
            elif isinstance(val, (int, float, datetime)):
                cleaned_row.append(val)
            else:
                num = parse_numeric(val)
                if num is not None:
                    cleaned_row.append(num)
                elif is_placeholder(val):
                    cleaned_row.append(None)
                else:
                    cleaned_row.append(val)

        data_rows.append([row_path, current_sec, sub_section] + cleaned_row)
        # Quy ước bảng tài chính: dòng 'Total ...' đầu tiên khép nhóm ('Assets:'); dòng sau nó không còn thuộc nhóm
        if major_section.endswith(':') and not is_indent_col and not use_indent and raw_label.lower().startswith('total'):
            major_section = ""

    final_cols = ["Row_Path", "Section", "Sub_Section"] + unique_col_names
    df = pd.DataFrame(data_rows, columns=final_cols)
    return df

def _mostly_code_columns(df: pd.DataFrame) -> bool:
    cols = [c for c in df.columns[3:] if not str(c).startswith(('Col_', 'Column_'))]
    return len(cols) >= 2 and sum(1 for c in cols if _CODE_RE.match(str(c))) / len(cols) >= 0.7


def _inherit_header(df: pd.DataFrame, base_df: Optional[pd.DataFrame]) -> pd.DataFrame:
    """Khối sau chỉ có mã cột -> dùng tên cột của khối đầu (cùng số cột) kèm mã để vẫn tra được theo tên."""
    if base_df is None or len(df.columns) != len(base_df.columns) or not _mostly_code_columns(df):
        return df
    names = list(df.columns[:3])
    for code, base in zip(df.columns[3:], base_df.columns[3:]):
        names.append(f"{base} | {code}" if _CODE_RE.match(str(code)) and not str(base).startswith(('Col_', 'Column_')) else code)
    out = df.copy()
    out.columns = names
    return out


def parse_workbook(file_path: str) -> Dict[str, pd.DataFrame]:
    """Parse toàn bộ workbook Excel, tự động tách các bảng độc lập trên cùng sheet."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheets_df = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        max_row = ws.max_row
        max_col = ws.max_column
        if max_row < 1 or max_col < 1:
            continue

        # Đọc toàn bộ ô vào grid và xử lý ô gộp (unmerge cells)
        grid = {}
        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                grid[(r, c)] = ws.cell(r, c).value

        for rng in ws.merged_cells.ranges:
            top_left_val = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    grid[(r, c)] = top_left_val

        table_ranges = detect_table_ranges(ws, grid, max_row, max_col)

        if len(table_ranges) == 1:
            start_r, end_r, _ = table_ranges[0]
            df = parse_table_block(ws, grid, start_r, end_r, max_col)
            if not df.empty:
                sheets_df[sheet_name] = df
        else:
            base_df = None
            for idx, (start_r, end_r, title) in enumerate(table_ranges):
                df = parse_table_block(ws, grid, start_r, end_r, max_col)
                if df.empty:
                    continue
                df = _inherit_header(df, base_df)
                if base_df is None and not _mostly_code_columns(df):
                    base_df = df
                clean_title = re.sub(r'[^\w\s-]', '', title).strip() or f"Table_{idx+1}"
                tbl_key = f"{sheet_name} - {clean_title}"
                if tbl_key in sheets_df:
                    tbl_key = f"{tbl_key} ({idx+1})"
                sheets_df[tbl_key] = df

    return sheets_df


def parse_raw_workbook(file_path: str) -> Dict[str, pd.DataFrame]:
    """Lưới ô nguyên bản (ô gộp đã điền, số dạng chuỗi đã chuyển thành số), chỉ số dòng = số dòng Excel.
    Đường lui cho LLM khi bảng đã parse sai cấu trúc."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    out = {}
    for ws in wb:
        if ws.max_row < 1 or ws.max_column < 1:
            continue
        grid = [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)] for r in range(1, ws.max_row + 1)]
        for rng in ws.merged_cells.ranges:
            val = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    grid[r - 1][c - 1] = val
        cleaned = []
        for row in grid:
            new = []
            for v in row:
                num = parse_numeric(v) if isinstance(v, str) else None
                new.append(num if num is not None else (None if is_placeholder(v) else v))
            cleaned.append(new)
        df = pd.DataFrame(cleaned, index=range(1, len(cleaned) + 1), columns=range(1, ws.max_column + 1))
        out[ws.title] = df
    return out


def suspicious_sheets(dfs: Dict[str, pd.DataFrame], raw_dfs: Dict[str, pd.DataFrame]) -> List[str]:
    """Tên các sheet mà bảng đã parse có dấu hiệu sai cấu trúc (không phụ thuộc nội dung cụ thể):
    ít dòng dữ liệu so với số dòng thực của sheet, tên cột chứa cả dữ liệu, hoặc phần lớn cột không có tên."""
    flagged = []
    for sheet, raw in raw_dfs.items():
        raw_rows = int(raw.notna().any(axis=1).sum())
        tables = [d for k, d in dfs.items() if k == sheet or k.startswith(f"{sheet} - ")]
        if not tables or raw_rows < 6:
            continue
        parsed_rows = sum(len(d) for d in tables)
        ncols = sum(max(d.shape[1] - 3, 0) for d in tables) or 1
        generic = sum(1 for d in tables for c in d.columns[3:] if re.fullmatch(r"Col_\d+", str(c)))
        wide_names = sum(1 for d in tables for c in d.columns[3:] if str(c).count(" | ") >= 6)
        if parsed_rows * 3 < raw_rows or wide_names or generic * 10 >= ncols * 4:
            flagged.append(sheet)
    return flagged


def raw_grid_preview(raw: pd.DataFrame, max_rows: int = 40, max_cols: int = 14, cell_chars: int = 24) -> str:
    """Bản xem lưới ô nguyên bản dạng văn bản (chỉ số dòng/cột = số dòng/cột Excel) cho sheet nghi ngờ."""
    lines = []
    for idx, row in raw.iloc[:, :max_cols].iterrows():
        cells = ["" if pd.isna(v) else str(v).replace("\n", " ").strip()[:cell_chars] for v in row.tolist()]
        if any(cells):
            lines.append(f"r{idx}: " + " | ".join(cells))
        if len(lines) >= max_rows:
            lines.append("...")
            break
    return "\n".join(lines)


def generate_table_summary(sheets_df: Dict[str, pd.DataFrame]) -> str:
    """Tạo tóm tắt cấu trúc bảng chi tiết kèm danh sách Row_Path cho LLM prompt."""
    summary_lines = []
    for sheet_name, df in sheets_df.items():
        summary_lines.append(f"### Table / Sheet: '{sheet_name}' (Shape: {df.shape[0]} rows x {df.shape[1]} columns)")
        summary_lines.append("Columns:")
        for col in df.columns:
            summary_lines.append(f"  - `{col}`")

        sections = [s for s in df['Section'].dropna().unique() if str(s).strip()]
        if sections:
            summary_lines.append(f"Major Sections: {sections[:10]}")

        # Cung cấp danh mục toàn bộ Row_Paths cho LLM tra cứu chính xác
        row_paths = [str(rp) for rp in df['Row_Path'].dropna().unique() if str(rp).strip()]
        summary_lines.append(f"Available Row_Paths ({len(row_paths)} rows):")
        # In toàn bộ nếu dưới 150 rows (đủ bao quát toàn bộ t07 có 114 rows mà không bị cắt)
        if len(row_paths) <= 150:
            for rp in row_paths:
                summary_lines.append(f"  * {rp}")
        else:
            for rp in row_paths[:120]:
                summary_lines.append(f"  * {rp}")
            summary_lines.append(f"  * ... ({len(row_paths) - 120} more rows)")

        summary_lines.append("Sample Data (first 2 rows):")
        summary_lines.append(df.head(2).to_string(index=False))
        summary_lines.append("")
    return "\n".join(summary_lines)
